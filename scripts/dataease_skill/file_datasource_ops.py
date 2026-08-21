from __future__ import annotations

import base64
import csv
import hashlib
import json
import mimetypes
from pathlib import Path
from typing import Any

from .audit import AuditLog
from .client import DataEaseClient
from .config import Settings
from .errors import DataEaseError
from .safety import PlanStore
from .versioning import adapter_for_client


MAX_FILE_BYTES = 500 * 1024 * 1024
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


def _envelope(operation: str, result: Any, **extra: Any) -> dict[str, Any]:
    return {"schema_version": 1, "ok": True, "operation": operation, "result": result,
            "warnings": extra.pop("warnings", []), **extra}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _fail(message: str, code: str = "invalid_file") -> None:
    raise DataEaseError(message, code=code, stage="file_datasource")


def inspect_file(file_name: str | Path) -> dict[str, Any]:
    """Validate a local CSV/XLSX without exposing its data in CLI output."""
    path = Path(file_name).expanduser().resolve()
    if not path.is_file():
        _fail(f"文件不存在或不是普通文件: {path}", "file_not_found")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        _fail("仅支持 .csv 或 .xlsx；旧版 .xls 请先另存为 .xlsx", "unsupported_file_type")
    size = path.stat().st_size
    if size == 0:
        _fail("文件为空")
    if size > MAX_FILE_BYTES:
        _fail("文件超过 DataEase Excel/CSV 上传上限 500 MiB", "file_too_large")

    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream)
            headers = next(reader, [])
        sheets = ["CSV"]
    else:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            _fail("处理 .xlsx 需要 openpyxl；请先 pip install -r requirements.txt", "dependency_missing")
            raise AssertionError from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = list(workbook.sheetnames)
        if not sheets:
            _fail("Excel 文件没有工作表")
        worksheet = workbook[sheets[0]]
        headers = ["" if cell.value is None else str(cell.value).strip() for cell in next(worksheet.iter_rows(max_row=1), ())]
        workbook.close()

    if not headers or any(not header for header in headers):
        _fail("首行必须是非空字段名，不能包含空表头")
    if len(set(headers)) != len(headers):
        _fail("首行字段名不能重复")
    return {"path": str(path), "format": suffix.lstrip("."), "bytes": size,
            "sha256": _sha256(path), "sheets": sheets, "columns": len(headers)}


def _load_generation_spec(spec_path: str) -> dict[str, Any]:
    raw = Path(spec_path).read_text(encoding="utf-8-sig") if spec_path != "-" else __import__("sys").stdin.read()
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise DataEaseError(f"JSON 配置无效: {exc}", code="invalid_spec", stage="input") from exc
    if not isinstance(spec, dict) or not isinstance(spec.get("columns"), list) or not isinstance(spec.get("rows"), list):
        raise DataEaseError("生成规格必须包含 columns 数组和 rows 数组", code="invalid_spec", stage="input")
    columns = [str(item.get("name") if isinstance(item, dict) else item).strip() for item in spec["columns"]]
    if not columns or any(not item for item in columns) or len(set(columns)) != len(columns):
        raise DataEaseError("columns 必须是唯一的非空字段名", code="invalid_spec", stage="input")
    normalized: list[list[Any]] = []
    for row in spec["rows"]:
        if isinstance(row, dict):
            normalized.append([row.get(column) for column in columns])
        elif isinstance(row, list) and len(row) == len(columns):
            normalized.append(row)
        else:
            raise DataEaseError("每一行必须是与 columns 等长的数组，或以字段名为键的对象", code="invalid_spec", stage="input")
    return {"columns": columns, "rows": normalized, "sheet": str(spec.get("sheet") or "数据")}


def generate_file(args: Any) -> dict[str, Any]:
    spec = _load_generation_spec(args.spec)
    output = Path(args.output).expanduser().resolve()
    suffix = output.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DataEaseError("--output 必须以 .csv 或 .xlsx 结尾", code="unsupported_file_type", stage="input")
    output.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".csv":
        with output.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(spec["columns"])
            writer.writerows(spec["rows"])
    else:
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise DataEaseError("生成 .xlsx 需要 openpyxl；请先 pip install -r requirements.txt", code="dependency_missing", stage="runtime") from exc
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = spec["sheet"][:31] or "数据"
        worksheet.append(spec["columns"])
        for row in spec["rows"]:
            worksheet.append(row)
        workbook.save(output)
    info = inspect_file(output)
    return _envelope("file-datasource.generate", {"file": info["path"], "format": info["format"],
        "rows": len(spec["rows"]), "columns": info["columns"], "sha256": info["sha256"]}, artifacts=[info["path"]])


def _context(client: DataEaseClient) -> dict[str, Any]:
    version = None
    try:
        version = client.data("GET", "/license/version")
    except DataEaseError:
        pass
    return {"base_url": client.settings.base_url, "api_prefix": client.settings.api_prefix,
            "org_id": client.settings.org_id or None, "version": version}


def _upload(client: DataEaseClient, info: dict[str, Any]) -> dict[str, Any]:
    path = Path(info["path"])
    mime_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    with path.open("rb") as stream:
        response = client.request("POST", "/datasource/uploadFile", files={"file": (path.name, stream, mime_type)},
                                  form={"id": "0", "editType": "0"})
    value = response.get("data") if isinstance(response, dict) and "data" in response else response
    if not isinstance(value, dict) or not isinstance(value.get("sheets"), list) or not value["sheets"]:
        raise DataEaseError("文件已上传，但 DataEase 未返回可创建数据源的工作表信息", code="invalid_upload_response", stage="upload")
    return value


def _source_payload(uploaded: dict[str, Any], name: str, pid: str, requested_sheets: list[str]) -> dict[str, Any]:
    sheets = uploaded.get("sheets") or []
    if requested_sheets:
        requested = set(requested_sheets)
        sheets = [sheet for sheet in sheets if str(sheet.get("sheetId")) in requested or str(sheet.get("tableName")) in requested]
    if not sheets:
        _fail("上传文件中未找到所选工作表", "sheet_not_found")
    cleaned: list[dict[str, Any]] = []
    for original in sheets:
        sheet = dict(original)
        sheet["data"] = []
        sheet["jsonArray"] = []
        fields = sheet.get("fields") if isinstance(sheet.get("fields"), list) else []
        if not any(field.get("checked", True) for field in fields if isinstance(field, dict)):
            _fail(f"工作表 {sheet.get('tableName') or sheet.get('sheetId')} 没有可用字段", "empty_sheet")
        cleaned.append(sheet)
    configuration = base64.b64encode(json.dumps(cleaned, ensure_ascii=False, separators=(",", ":")).encode("utf-8")).decode("ascii")
    return {"name": name, "pid": pid, "type": "Excel", "sheets": cleaned, "editType": 0, "configuration": configuration}


def _public(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        return {}
    return {key: value.get(key) for key in ("id", "name", "pid", "type", "status", "nodeType") if key in value}


def create_file_datasource(args: Any, settings: Settings, client: DataEaseClient, plans: PlanStore, audit: AuditLog) -> dict[str, Any]:
    operation = "file-datasource.create"
    info = inspect_file(args.file)
    target = {"file_sha256": info["sha256"], "name": str(args.name or Path(info["path"]).stem), "pid": str(args.pid or "0"),
              "sheets": list(args.sheet or [])}
    if not args.apply:
        plan = plans.create(operation, target=target,
            changes=[{"action": "upload-and-create", "resource": "Excel datasource", "format": info["format"],
                      "bytes": info["bytes"], "sheets": info["sheets"], "columns": info["columns"]}],
            risk="L1", spec={"file_sha256": info["sha256"], "name": target["name"], "pid": target["pid"], "sheets": target["sheets"]},
            rollback={"strategy": "delete-created-datasource", "requires_confirmation": True}, context=_context(client))
        return _envelope(operation, plan, mode="dry-run", changes=plan["changes"], warnings=[
            "dry-run 只校验本地文件，不上传任何数据。apply 会上传文件并创建 Excel 数据源。",
            "创建后请用 dataset quick-create 生成数据集，再用 visual autopilot 生成仪表板或 DataV。",
        ])
    if not args.plan_id:
        raise DataEaseError("执行创建需要 --plan-id", code="plan_required", stage="safety")
    version, adapter = adapter_for_client(client)
    adapter.require("file_datasource", version, mutation=True, client=client)
    plan = plans.load(args.plan_id, getattr(args, "confirm_token", ""), expected_context=_context(client))
    spec = plan.get("spec", {})
    if plan.get("operation") != operation or spec.get("file_sha256") != info["sha256"] or spec.get("name") != target["name"] or spec.get("pid") != target["pid"] or spec.get("sheets") != target["sheets"]:
        raise DataEaseError("文件、名称、目录或工作表选择与 dry-run 不一致", code="spec_changed", stage="safety")
    uploaded = _upload(client, info)
    saved = client.data("POST", "/datasource/save", _source_payload(uploaded, target["name"], target["pid"], target["sheets"]))
    datasource_id = str(saved.get("id") if isinstance(saved, dict) else saved or "")
    if not datasource_id:
        raise DataEaseError("DataEase 未返回新数据源 ID，无法完成回读验证", code="missing_resource_id", stage="verification")
    detail = client.data("GET", f"/datasource/hidePw/{datasource_id}")
    public = _public(detail if isinstance(detail, dict) else saved)
    if str(public.get("name")) != target["name"] or str(public.get("type")) != "Excel":
        raise DataEaseError("创建后的数据源回读与请求不一致", code="verification_failed", stage="verification")
    audit_id = audit.write(operation, status="success", risk="L1", target=plan.get("target"), changes=plan.get("changes"), result=public)
    plans.mark_applied(args.plan_id, audit_id)
    return _envelope(operation, public, mode="apply", adapter="official-api", audit_id=audit_id, artifacts=[info["path"]], warnings=[
        "数据源已创建。下一步请先为数据源中的目标表执行 dataset quick-create dry-run，再确认创建数据集。",
    ])
